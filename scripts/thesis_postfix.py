"""Post-build fixes for the thesis .docx (terminology, front-matter structure, appendix placeholders).

Usage: python scripts/thesis_postfix.py <docx> [--commit <sha>] [--login-user U --login-password P]

The chapters 1-2 text lives in the Word template, so terminology fixes that are already applied to the
Markdown chapters are applied here too. Placeholders for the demo login are filled only when the
values are passed on the command line (they are never stored in the repository).
"""
from __future__ import annotations

import argparse
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

TERMS = [
    ("biên nhận tính toán", "bản ghi truy vết phép tính"),
    ("biên nhận tính", "bản ghi truy vết phép tính"),
    ("ngăn kéo nguồn", "khung trượt tra cứu nguồn"),
    ("hợp đồng dữ liệu", "đặc tả ràng buộc dữ liệu"),
    ("hợp đồng đầu ra", "đặc tả định dạng đầu ra"),
    ("hợp đồng giao diện", "đặc tả giao diện"),
    ("suy giảm an toàn", "cơ chế dự phòng an toàn"),
    ("Bộ dữ liệu khóa", "Bộ dữ liệu đã khóa (cố định)"),
    ("bộ dữ liệu khóa", "bộ dữ liệu đã khóa (cố định)"),
    ("Lắp báo cáo", "Tổng hợp báo cáo"),
    ("lắp báo cáo", "tổng hợp báo cáo"),
    ("Lời nhắc", "Prompt"),
    ("lời nhắc", "prompt"),
    ("cũng cố", "củng cố"),
    ("bằng công thức có phiên bản", "theo công thức được định danh phiên bản"),
    ("công thức có phiên bản", "công thức được định danh phiên bản"),
]

CONVENTION = (" Trong khóa luận này, \"tác tử\" là cách gọi chung của agent; trong ngữ cảnh nghiệp vụ của hệ thống, "
              "các tác tử chuyên trách được gọi là \"chuyên viên\" (chuyên viên doanh nghiệp, chuyên viên ngành).")


def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    yield p


def replace_in_paragraph(p, old, new) -> bool:
    if old not in p.text:
        return False
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    # term split across runs: collapse into the first run (formatting of this paragraph is simplified)
    text = p.text.replace(old, new)
    for r in p.runs[1:]:
        r.text = ""
    p.runs[0].text = text
    return True


def set_page_break_before(p, on=True):
    p.paragraph_format.page_break_before = on


def drop_page_break_runs(p):
    for br in p._p.findall(".//" + qn("w:br")):
        if br.get(qn("w:type")) == "page":
            run = br.getparent()
            run.remove(br)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--commit", default=None)
    ap.add_argument("--login-user", default=None)
    ap.add_argument("--login-password", default=None)
    a = ap.parse_args()
    doc = Document(a.docx)

    # 1. terminology
    n = 0
    for p in all_paragraphs(doc):
        for old, new in TERMS:
            while replace_in_paragraph(p, old, new):
                n += 1
    print("terminology fixes:", n)

    # 2. convention sentence after the first "tác tử (agent)" definition in chapter 1
    for p in doc.paragraphs:
        if "tác tử (agent)" in p.text and "Trong khóa luận này" not in p.text:
            p.runs[-1].text = p.runs[-1].text.rstrip() + CONVENTION
            print("convention sentence added")
            break

    # 3. front matter: list headings are not TOC headings; chapter headings start a new page
    front = {"MỤC LỤC", "DANH MỤC HÌNH ẢNH", "DANH MỤC BẢNG BIỂU"}
    for p in doc.paragraphs:
        t = p.text.strip()
        if t in front and p.style.name.startswith("Heading"):
            drop_page_break_runs(p)
            p.style = doc.styles["Normal"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
            set_page_break_before(p, True)
        elif p.style.name == "Heading 1" and t.upper().startswith("CHƯƠNG"):
            drop_page_break_runs(p)
            set_page_break_before(p, True)
        elif p.style.name == "Heading 1" and t.upper().startswith(("TÀI LIỆU THAM KHẢO", "PHỤ LỤC")):
            drop_page_break_runs(p)
            set_page_break_before(p, True)

    # 4. placeholders
    filled = 0
    for p in all_paragraphs(doc):
        t = p.text
        if "ĐIỀN" not in t:
            continue
        new = t
        if "Đăng nhập" in t and a.login_user and a.login_password:
            new = new.replace("[ĐIỀN]", a.login_user, 1).replace("[ĐIỀN]", a.login_password, 1)
        if "MÃ COMMIT" in t and a.commit:
            new = new.replace("[ĐIỀN: MÃ COMMIT]", a.commit)
        if "LIÊN KẾT TẢI" in t:
            new = new.replace("tải tại `[ĐIỀN: LIÊN KẾT TẢI]`", "cung cấp kèm hồ sơ nộp bài hoặc theo yêu cầu")
            new = new.replace("tải tại [ĐIỀN: LIÊN KẾT TẢI]", "cung cấp kèm hồ sơ nộp bài hoặc theo yêu cầu")
        if new != t:
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = new
            filled += 1
    print("placeholders filled:", filled)
    doc.save(a.docx)
    left = [p.text[:60] for p in all_paragraphs(Document(a.docx)) if "ĐIỀN" in p.text]
    print("placeholders left:", len(left))


if __name__ == "__main__":
    main()

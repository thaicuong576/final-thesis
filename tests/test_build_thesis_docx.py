"""
Tests for scripts/build_thesis_docx.py

Runs a minimal build from fixture Markdown content and a programmatically
created minimal .docx template (no external files needed).
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
import unittest
from pathlib import Path

# Make sure scripts/ is importable
REPO_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import build_thesis_docx as bld
import thesis_diagrams as diagrams

# ---------------------------------------------------------------------------
# Fixture helpers
# ---------------------------------------------------------------------------

MINIMAL_MARKDOWN = """\
# CHƯƠNG 1. GIỚI THIỆU

Đây là đoạn văn bình thường với **chữ đậm** và *chữ nghiêng* và `code`.

## 1.1. Bối cảnh

Đoạn giới thiệu bối cảnh nghiên cứu.

### 1.1.1. Chi tiết

Nội dung chi tiết đơn vị nhỏ nhất.

- Mục 1
- Mục 2 **đậm**

1. Bước đầu tiên
2. Bước thứ hai

| Cột A | Cột B | Cột C |
|-------|-------|-------|
| **Dữ liệu** | 100 | 200 |
| Giá trị | 300 | 400 |

Bảng 1.1: Bảng ví dụ minh họa kết quả.

> Đây là đoạn blockquote trích dẫn tài liệu.

Đây là đoạn văn cuối chương có nội dung bình thường.
"""

FRONT_MATTER_MD = """\
# TRƯỜNG ĐẠI HỌC KINH TẾ

**Khóa luận tốt nghiệp**

Sinh viên: [ĐIỀN]
"""


def _make_minimal_template(path: str) -> None:
    """Create a minimal .docx that serves as the template for tests."""
    doc = Document()
    # Set page dimensions like the real template
    sec = doc.sections[0]
    sec.page_width = Cm(21.59)
    sec.page_height = Cm(27.94)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.54)
    sec.right_margin = Cm(2.54)

    # Add a dummy paragraph so the doc is not empty
    doc.add_paragraph("TEMPLATE CONTENT TO BE CLEARED")
    doc.save(path)


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestPrepareTemplate(unittest.TestCase):
    """Test that prepare_template clears body content but keeps sectPr."""

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp()
        self.template_path = os.path.join(self.tmp_dir, "template.docx")
        _make_minimal_template(self.template_path)

    def test_template_cleared(self):
        doc = bld.prepare_template(self.template_path)
        # Body should be empty (only sectPr remains)
        body = doc.element.body
        non_secPr = [el for el in body if el.tag != qn("w:sectPr")]
        self.assertEqual(len(non_secPr), 0, "Body should have no elements besides sectPr after clearing")

    def test_sectPr_preserved(self):
        doc = bld.prepare_template(self.template_path)
        body = doc.element.body
        sectPr = body.find(qn("w:sectPr"))
        self.assertIsNotNone(sectPr, "sectPr must be preserved in the template")


class TestConvertMarkdown(unittest.TestCase):
    """Test Markdown → Word conversion."""

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp()
        self.template_path = os.path.join(self.tmp_dir, "template.docx")
        _make_minimal_template(self.template_path)
        self.doc = bld.prepare_template(self.template_path)
        bld.convert_markdown_to_doc(self.doc, MINIMAL_MARKDOWN, is_first_chapter=True)

    def _paragraphs_by_style(self, style_name: str):
        return [p for p in self.doc.paragraphs if p.style and p.style.name == style_name]

    def test_heading1_style(self):
        h1s = self._paragraphs_by_style("Heading 1")
        self.assertGreater(len(h1s), 0, "Should have at least one Heading 1")

    def test_heading2_style(self):
        h2s = self._paragraphs_by_style("Heading 2")
        self.assertGreater(len(h2s), 0, "Should have at least one Heading 2")

    def test_heading3_style(self):
        h3s = self._paragraphs_by_style("Heading 3")
        self.assertGreater(len(h3s), 0, "Should have at least one Heading 3")

    def test_normal_paragraphs_exist(self):
        normals = self._paragraphs_by_style("Normal")
        self.assertGreater(len(normals), 0, "Should have Normal paragraphs")

    def test_heading1_is_uppercase(self):
        h1s = self._paragraphs_by_style("Heading 1")
        for h1 in h1s:
            text = h1.text
            if text:
                self.assertEqual(text, text.upper(), f"Heading 1 should be uppercase: {text!r}")

    def test_heading1_font_size(self):
        h1s = self._paragraphs_by_style("Heading 1")
        self.assertGreater(len(h1s), 0)
        for h1 in h1s:
            for run in h1.runs:
                if run.font.size is not None:
                    self.assertEqual(run.font.size, Pt(13), f"Heading 1 run should be 13pt, got {run.font.size}")

    def test_bold_run_in_normal(self):
        normals = self._paragraphs_by_style("Normal")
        bold_found = False
        for para in normals:
            for run in para.runs:
                if run.font.bold and run.text:
                    bold_found = True
                    break
        self.assertTrue(bold_found, "Should find at least one bold run in Normal paragraphs")

    def test_no_raw_markdown_in_text(self):
        """No paragraph text should contain raw Markdown syntax."""
        bad_patterns = ['**', '```', '<br', '<b>', '</b>']
        for para in self.doc.paragraphs:
            full_text = para.text
            for pat in bad_patterns:
                self.assertNotIn(pat, full_text,
                                 f"Raw Markdown/HTML found in paragraph text: {full_text!r}")

    def test_table_created(self):
        self.assertGreater(len(self.doc.tables), 0, "Should have at least one table")

    def test_table_header_bold(self):
        for table in self.doc.tables:
            if not table.rows:
                continue
            header_row = table.rows[0]
            for cell in header_row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.text.strip():
                            self.assertTrue(
                                run.font.bold,
                                f"Table header cell should be bold: {run.text!r}"
                            )

    def test_caption_style(self):
        caption_paras = [p for p in self.doc.paragraphs if bld._is_caption_line(p.text)]
        self.assertGreater(len(caption_paras), 0, "Should have at least one caption paragraph")
        for cap in caption_paras:
            # Caption style OR italic runs
            is_caption_style = (cap.style and "Caption" in cap.style.name)
            has_italic = any(run.font.italic for run in cap.runs if run.text.strip())
            self.assertTrue(
                is_caption_style or has_italic,
                f"Caption paragraph should have Caption style or italic: {cap.text!r}"
            )


class TestBuildFull(unittest.TestCase):
    """Integration test: build → save → reopen → verify."""

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp()
        self.template_path = os.path.join(self.tmp_dir, "template.docx")
        self.out_path = os.path.join(self.tmp_dir, "out.docx")
        # Create front-matter and chapter files in tmp
        self.fm_path = os.path.join(self.tmp_dir, "front-matter.md")
        self.ch_path = os.path.join(self.tmp_dir, "chapter.md")
        _make_minimal_template(self.template_path)
        Path(self.fm_path).write_text(FRONT_MATTER_MD, encoding="utf-8")
        Path(self.ch_path).write_text(MINIMAL_MARKDOWN, encoding="utf-8")

    def test_build_produces_docx(self):
        doc = bld.prepare_template(self.template_path)
        fm_text = Path(self.fm_path).read_text(encoding="utf-8")
        bld.add_front_matter(doc, fm_text)
        ch_text = Path(self.ch_path).read_text(encoding="utf-8")
        bld.convert_markdown_to_doc(doc, ch_text, is_first_chapter=True)
        doc.save(self.out_path)

        self.assertTrue(os.path.exists(self.out_path), "Output docx must be created")
        self.assertGreater(os.path.getsize(self.out_path), 1000, "Output file seems too small")

    def test_reopen_and_check_content(self):
        doc = bld.prepare_template(self.template_path)
        ch_text = Path(self.ch_path).read_text(encoding="utf-8")
        bld.convert_markdown_to_doc(doc, ch_text, is_first_chapter=True)
        doc.save(self.out_path)

        doc2 = Document(self.out_path)
        h1s = [p for p in doc2.paragraphs if p.style and p.style.name == "Heading 1"]
        self.assertGreater(len(h1s), 0, "Reopened doc should have Heading 1 paragraphs")

    def test_font_tnr_on_normal(self):
        doc = bld.prepare_template(self.template_path)
        bld.convert_markdown_to_doc(doc, MINIMAL_MARKDOWN, is_first_chapter=True)
        doc.save(self.out_path)
        doc2 = Document(self.out_path)
        normals = [p for p in doc2.paragraphs if p.style and p.style.name == "Normal"]
        for para in normals[:5]:
            for run in para.runs:
                if run.text.strip() and not run.font.name:
                    # Font can be inherited - if set, must be TNR
                    pass
                if run.font.name:
                    self.assertIn(
                        run.font.name, ["Times New Roman", "Consolas"],
                        f"Run font should be Times New Roman or Consolas, got {run.font.name!r}"
                    )


class TestInlineMarkdown(unittest.TestCase):
    """Test inline parsing utilities."""

    def _make_doc_with_para(self):
        tmp = tempfile.mktemp(suffix=".docx")
        doc = Document()
        return doc

    def test_clean_html_br(self):
        result = bld._clean_html("line1<br/>line2")
        self.assertIn("\n", result)
        self.assertNotIn("<br", result)

    def test_clean_html_bold(self):
        result = bld._clean_html("<b>bold</b>")
        self.assertEqual(result, "**bold**")

    def test_clean_html_strips_remaining(self):
        result = bld._clean_html("<span>text</span>")
        self.assertEqual(result, "text")

    def test_no_raw_markdown_after_inline(self):
        doc = Document()
        para = doc.add_paragraph()
        bld._add_inline_runs(para, "**bold** and *italic* and `code`")
        full_text = "".join(r.text for r in para.runs)
        for pat in ['**', '```', '*']:
            self.assertNotIn(pat, full_text, f"Raw Markdown {pat!r} leaked into run text")

    def test_is_caption_hinh(self):
        self.assertTrue(bld._is_caption_line("Hình 3.1: Quy trình"))
        self.assertFalse(bld._is_caption_line("Đây là đoạn bình thường"))

    def test_is_caption_bang(self):
        self.assertTrue(bld._is_caption_line("Bảng 2.5: Dữ liệu tổng hợp"))


class TestMermaidParser(unittest.TestCase):
    """Test the offline Mermaid parser in scripts/thesis_diagrams.py."""

    def test_parse_chain_groups_labels_and_inline_nodes(self):
        g = diagrams.parse_mermaid("""flowchart TD
    subgraph S ["Nhóm"]
        A["<b>Bắt đầu</b><br/>• chi tiết"] --> B[Xử lý] --> C[Kết thúc]
    end
    A & B -.->|"Phát hiện lỗi (≤ 2 vòng)"| D(["Kết quả"])
""")
        self.assertEqual(g.labels["A"], "Bắt đầu" + chr(10) + "• chi tiết")
        self.assertIn(("A", "B", "-->", ""), g.edges)
        self.assertIn(("B", "C", "-->", ""), g.edges)
        self.assertIn(("A", "D", "-.->", "Phát hiện lỗi (≤ 2 vòng)"), g.edges)
        self.assertIn(("B", "D", "-.->", "Phát hiện lỗi (≤ 2 vòng)"), g.edges)
        self.assertNotIn("lỗi", g.labels)
        self.assertEqual(g.groups[0][0], "Nhóm")
        self.assertEqual(set(g.groups[0][1]), {"A", "B", "C"})

    def test_parse_empty(self):
        g = diagrams.parse_mermaid("")
        self.assertEqual(g.edges, [])

    def test_render_writes_png(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "x.mmd"
            src.write_text("flowchart LR" + chr(10) + "    A[Một] --> B[Hai]" + chr(10), encoding="utf-8")
            out = diagrams.render(src, Path(d) / "x.png")
            self.assertTrue(out.exists() and out.stat().st_size > 0)


if __name__ == "__main__":
    unittest.main()

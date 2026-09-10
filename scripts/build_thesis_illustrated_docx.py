"""Add evidence-linked native tables and verified figures to the thesis."""
from pathlib import Path
import hashlib
import json
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
import build_thesis_essay_docx as base

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'docs/assets/thesis-visuals-20260910'
QA = ROOT / 'reports/thesis-visuals-20260910'
SOURCE = ROOT / 'reports/analysis-report-v2-final-evaluation-gold-v5-20260909-v4/final-evaluation.json'
ARMS = ['single_agent','hierarchical_no_critique','hierarchical_bounded_critique']
ARM_NAMES = ['Một tác tử','Phân cấp không phản biện','Phân cấp có phản biện']


def number(value, digits=6):
    return f'{value:.{digits}f}'.rstrip('0').rstrip('.').replace('.', ',')


def main():
    base.main()
    doc = Document(base.OUTPUT)
    original = [p.text for p in doc.paragraphs]
    for name in ('Caption', 'Table Source'):
        if name not in doc.styles:
            doc.styles.add_style(name, 1)
        style = doc.styles[name]
        base.configure_font(style)
        style.font.bold = False
        style.font.italic = False
        pf = style.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.widow_control = True
    evaluation = json.loads(SOURCE.read_text(encoding='utf-8'))
    summary = {(r['issuer_count'],r['arm']):r for r in evaluation['quality']['arm_summaries']}
    audit = json.loads((ROOT / 'reports/data-foundation/gold-bank-2025-expanded-v5-evidence.json').read_text(encoding='utf-8'))
    inventory = []
    cursor = None

    def anchor(title):
        nonlocal cursor
        paragraphs = doc.paragraphs
        matches = [i for i,p in enumerate(paragraphs) if p.text == title and p.style.name.startswith('Heading')]
        assert len(matches) == 1, title
        end = matches[0]
        for i in range(end+1, len(paragraphs)):
            if paragraphs[i].style.name.startswith('Heading'):
                break
            end = i
        cursor = paragraphs[end]._p

    def paragraph(text, style='Normal', keep=False):
        nonlocal cursor
        p = doc.add_paragraph(text, style)
        cursor.addnext(p._p)
        cursor = p._p
        p.paragraph_format.keep_with_next = keep
        return p

    def figure(filename, caption, source, lead, alt):
        nonlocal cursor
        paragraph(lead, keep=True)
        p = paragraph('', keep=True)
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.space_after = Pt(6)
        picture = p.add_run().add_picture(str(ASSETS / filename), width=Inches(6.5))
        picture._inline.docPr.set('descr', alt)
        cp = paragraph(caption, 'Caption', keep=True)
        cp.runs[0].bold = True
        paragraph(source, 'Table Source')
        inventory.append({'type':'figure','caption':caption,'file':filename,
                          'sha256':hashlib.sha256((ASSETS/filename).read_bytes()).hexdigest()})

    def table(caption, headers, rows, widths, source, lead, numeric=True):
        nonlocal cursor
        paragraph(lead, keep=True)
        cp = paragraph(caption,'Caption',keep=True)
        cp.runs[0].bold = True
        tab = doc.add_table(rows=1, cols=len(headers))
        cursor.addnext(tab._tbl)
        cursor = tab._tbl
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        tab.autofit = False
        for col,width in zip(tab.columns,widths):
            col.width = Inches(width)
        props = tab._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            item = OxmlElement('w:'+edge)
            for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]:
                item.set(qn('w:'+key),val)
            borders.append(item)
        props.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for edge in ['top','left','bottom','right']:
            item = OxmlElement('w:'+edge)
            item.set(qn('w:w'),'85')
            item.set(qn('w:type'),'dxa')
            margins.append(item)
        props.append(margins)
        for values in rows:
            tab.add_row()
        for ridx, values in enumerate([headers]+rows):
            row = tab.rows[ridx]
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            if ridx == 0:
                trpr.append(OxmlElement('w:tblHeader'))
            for cidx,(cell,value,width) in enumerate(zip(row.cells,values,widths)):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = str(value)
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.5
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ridx == 0 or (numeric and cidx>0) else WD_ALIGN_PARAGRAPH.LEFT
                # Keep the short resource comparison on one page; larger tables
                # may flow across pages with a repeated header and intact rows.
                p.paragraph_format.keep_with_next = ridx == 0 or (
                    caption.startswith('Bảng 6.3') and ridx < len(rows)
                )
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(13)
                    run.font.color.rgb = RGBColor(0,0,0)
                    run.bold = ridx == 0
                if ridx == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'),'EEEEEE')
                    cell._tc.get_or_add_tcPr().append(shade)
        paragraph(source,'Table Source')
        inventory.append({'type':'table','caption':caption,'headers':headers,'rows':rows})

    anchor('Lựa chọn mẫu và hình thành dữ liệu chuẩn')
    rows=[]
    for issuer in ['ACB','TCB','VPB','MBB','VCB']:
        count=audit['by_issuer'][issuer]['valid_observations']
        rows.append([issuer,16,count,16-count,number(count/16*100,1)+'%'])
    rows.append(['Tổng',80,76,4,'95%'])
    table('Bảng 3.1 Phạm vi và độ đầy đủ của dữ liệu Gold',
          ['Ngân hàng','Mục tiêu','Hợp lệ','Thiếu','Độ phủ số'],rows,[1.35,1.1,1.1,1.1,1.85],
          'Nguồn: Báo cáo kiểm tra Gold v5, ngày 9 tháng 9 năm 2026. Bốn vị trí thiếu đều thuộc ACB quý I năm 2024. Độ phủ số khác với độ phủ bằng chứng trên giá trị hợp lệ.',
          'Bảng 3.1 trình bày phạm vi dữ liệu theo từng ngân hàng. Việc giữ riêng cột thiếu cho thấy bộ dữ liệu không có cùng mức độ đầy đủ ở mọi thành viên của mẫu.')

    anchor('Phân chia trách nhiệm xử lý')
    figure('figure-4-1-architecture.png','Hình 4.1 Kiến trúc từ nguồn dữ liệu đến báo cáo',
           'Nguồn: Tổng hợp từ kiến trúc và mã triển khai của hệ thống. Sơ đồ minh họa bằng GPT Image Gen, được đối chiếu lại với luồng xử lý.',
           'Hình 4.1 phân biệt khâu chuẩn bị dữ liệu với khâu phân tích. Tác tử chỉ sử dụng dữ liệu đã chấp nhận, còn giao diện đọc kết quả sau bước kiểm tra đầu ra.',
           'Nguồn công bố đi qua chuẩn hóa, đối chiếu và Gold. Công cụ chỉ đọc cấp dữ liệu cho tác tử. Đầu ra qua kiểm tra, JSON, Markdown và HTML rồi giao diện. Nhật ký được lưu riêng.')

    anchor('Phạm vi truy vấn và công cụ của tác tử')
    table('Bảng 4.1 Vai trò và phạm vi truy cập trong luồng phân tích',
          ['Vai trò','Trách nhiệm','Công cụ trong phạm vi ngân hàng'],[
            ['Một tác tử','Xử lý yêu cầu và tổng hợp','Chỉ tiêu, nhóm so sánh, bằng chứng và trạng thái dữ liệu'],
            ['Công ty','Phân tích ngân hàng và diễn biến','Chỉ tiêu công ty, bằng chứng và trạng thái dữ liệu'],
            ['Nhóm so sánh','Đối chiếu các ngân hàng','Chuẩn nhóm, bằng chứng và trạng thái dữ liệu'],
            ['Điều phối','Tổng hợp kết quả của worker','Bằng chứng, thông tin nguồn và trạng thái dữ liệu'],
            ['Kiểm chứng','Kiểm tra nội dung được yêu cầu','Chỉ tiêu, chuẩn nhóm, bằng chứng và trạng thái dữ liệu']
          ],[1.25,2.0,3.25],
          'Nguồn: Các bộ công cụ theo vai trò trong agent_tools.py và luồng composite_runner.py. Tất cả truy cập chỉ đọc, bị giới hạn bởi yêu cầu. Công cụ vĩ mô có trong mã chung nhưng không thuộc phạm vi bảng này.',
          'Bảng 4.1 làm rõ việc phân công không chỉ là đổi tên tác tử. Bộ điều phối sử dụng kết quả của các thành phần phân tích và không được cấp công cụ truy vấn trực tiếp mọi chỉ tiêu như tác tử tổng hợp.',numeric=False)

    anchor('Kiểm tra đầu ra và phản biện')
    figure('figure-4-2-workflow.png','Hình 4.2 Luồng xử lý của ba nhánh thực nghiệm',
           'Nguồn: Luồng composite_runner.py tại phiên bản được kiểm tra. Sơ đồ minh họa bằng GPT Image Gen. Kiểm chứng được kích hoạt theo cấu hình và giới hạn tối đa hai lượt.',
           'Hình 4.2 mô tả ba cách tổ chức xử lý trên cùng yêu cầu. Kết quả kiểm chứng được lưu riêng, nên sơ đồ không có mũi tên quay lại tổng hợp. Thứ tự trong hình không khẳng định mọi worker chạy song song.',
           'Yêu cầu và Gold qua kiểm tra phạm vi rồi chọn một tác tử, phân cấp, hoặc phân cấp có kiểm chứng. Ba nhánh được kiểm tra và ghi nhật ký. Kiểm chứng không tổng hợp lại báo cáo.')

    anchor('Giao diện và luồng đọc báo cáo')
    figure('figure-5-1-overview.png','Hình 5.1 Màn hình tổng quan báo cáo trên ứng dụng local',
           'Nguồn: Ảnh chụp trực tiếp ứng dụng local ngày 10 tháng 9 năm 2026, sử dụng Gold v5. Đây là giao diện chỉ đọc báo cáo đã tạo, không phải một lần gọi mô hình mới.',
           'Hình 5.1 cho thấy thông tin phạm vi và độ đầy đủ dữ liệu được đặt cạnh nhận định tổng hợp. Các số 76/80 và 76/76 có mẫu số khác nhau, giúp người dùng phân biệt độ phủ số liệu với độ phủ bằng chứng.',
           'Ảnh thật màn hình tổng quan báo cáo ngân hàng, hiển thị Gold v5, 76 trên 80 vị trí hợp lệ, 4 vị trí thiếu và 76 trên 76 giá trị có bằng chứng.')
    figure('figure-5-2-evidence.png','Hình 5.2 Nhận định được mở cùng dấu vết bằng chứng',
           'Nguồn: Ảnh chụp trực tiếp một nhận định trong ứng dụng local ngày 10 tháng 9 năm 2026. Vùng chụp giữ nguyên nội dung giao diện, gồm tham chiếu dữ liệu và vị trí nguồn.',
           'Hình 5.2 minh họa cách người đọc kiểm tra nhận định về NPL của VPB. Phần mở rộng tách nhận định, diễn giải và giới hạn, sau đó trình bày dữ liệu tham chiếu. Đường dẫn trang là căn cứ để kiểm tra, không tự thay thế việc đọc tài liệu gốc.',
           'Ảnh thật nhận định VPB có NPL cao nhất quý IV năm 2024, cùng diễn giải, giới hạn và định vị trang nguồn của các số liệu tham chiếu.')

    anchor('Cơ sở số liệu và khả năng hoàn thành')
    heads=['Số ngân hàng']+ARM_NAMES
    rows=[[n]+[number(summary[n,a]['quality_mean']) for a in ARMS] for n in range(1,6)]
    table('Bảng 6.1 Chất lượng trung bình của các báo cáo hoàn thành',heads,rows,[1.0,1.4,2.05,2.05],
          'Nguồn: Kết quả đánh giá đóng băng ngày 9 tháng 9 năm 2026, bản v4. Điểm chuẩn hóa từ 0 đến 1, chỉ tính báo cáo hoàn thành và được chấm. Không có khoảng tin cậy trong bảng này.',
          'Bảng 6.1 tổng hợp chất lượng trên các báo cáo thực sự hoàn thành. Cần đọc đồng thời với Bảng 6.2, bởi điểm cao ở một cấu hình có ít đầu ra không đồng nghĩa với hiệu quả vận hành cao.')
    rows=[[n]+[f"{summary[n,a]['completed_repetitions']}/5" for a in ARMS] for n in range(1,6)]
    rows.append(['Tổng']+[f"{sum(summary[n,a]['completed_repetitions'] for n in range(1,6))}/25" for a in ARMS])
    table('Bảng 6.2 Số công việc hoàn thành theo cấu hình',heads,rows,[1.0,1.4,2.05,2.05],
          'Nguồn: Cùng đợt thực nghiệm với Bảng 6.1. Tổng cộng 61/75 công việc hoàn thành và 14 công việc thất bại. Mỗi ô theo cấp có mẫu số năm lượt dự kiến.',
          'Bảng 6.2 cho thấy nhánh có phản biện chỉ hoàn thành một trong năm lượt ở cấp năm ngân hàng. Điều này giải thích vì sao điểm chất lượng bằng 1 ở cấp đó chưa đủ để chọn kiến trúc.')

    anchor('Đánh đổi về chi phí và độ trễ ở ba ngân hàng')
    rows=[]
    for arm,name in zip(ARMS,ARM_NAMES):
        r=summary[3,arm]
        rows.append([name,number(r['cost_mean_usd']),number(r['latency_mean_ms']/1000,1),
                     f"{r['input_tokens_mean']:,.1f}".replace(',','X').replace('.',',').replace('X','.'),
                     f"{r['output_tokens_mean']:,.1f}".replace(',','X').replace('.',',').replace('X','.')])
    table('Bảng 6.3 Tài nguyên trung bình ở yêu cầu ba ngân hàng',
          ['Nhánh','USD mỗi lượt','Giây mỗi lượt','Token vào','Token ra'],rows,[1.7,1.2,1.0,1.3,1.3],
          'Nguồn: Nhật ký tổng hợp của đợt đánh giá cuối. Chỉ tính lượt hoàn thành, lần lượt 5, 5 và 2 lượt. Chi phí là ước tính theo usage, không phải tổng hóa đơn bao gồm mọi lượt lỗi. Token có phần thập phân vì là giá trị trung bình.',
          'Bảng 6.3 đặt tài nguyên cạnh điểm chuyển lợi thế đã quan sát. Chi phí và thời gian tăng cần được xem cùng mức cải thiện chất lượng, thay vì suy ra rằng đa tác tử luôn đáng sử dụng.')

    assert len([i for i in inventory if i['type']=='table']) == 5
    assert len([i for i in inventory if i['type']=='figure']) == 4
    doc.save(base.OUTPUT)
    markdown=[]
    figure_index=0
    figure_files=[i['file'] for i in inventory if i['type']=='figure']
    for element in doc.element.body:
        if element.tag==qn('w:p'):
            p=Paragraph(element,doc)
            if element.xpath('.//w:drawing'):
                markdown.append(f'![Hình minh họa](assets/thesis-visuals-20260910/{figure_files[figure_index]})')
                figure_index+=1
            elif p.text:
                prefix={'Title':'# ','Heading 1':'## ','Heading 2':'### '}.get(p.style.name,'')
                markdown.append(prefix+p.text)
        elif element.tag==qn('w:tbl'):
            tbl=Table(element,doc)
            lines=['| '+' | '.join(c.text.replace('\n',' ') for c in row.cells)+' |' for row in tbl.rows]
            lines.insert(1,'| '+' | '.join(['---']*len(tbl.columns))+' |')
            markdown.append('\n'.join(lines))
    (ROOT/'docs/THESIS-DRAFT-ILLUSTRATED-VI.md').write_text('\n\n'.join(markdown)+'\n',encoding='utf-8')
    check=Document(base.OUTPUT)
    assert len(check.tables)==5 and len(check.inline_shapes)==4
    # All original prose remains, in order, with added material between sections.
    it=iter([p.text for p in check.paragraphs])
    for text in original:
        assert any(t==text for t in it), text[:80]
    with ZipFile(base.OUTPUT) as z:
        assert not any(n.startswith(('word/header','word/footer')) for n in z.namelist())
        assert b'<w:pBdr' not in z.read('word/document.xml')
    for tab in check.tables:
        for row in tab.rows:
            for cell in row.cells:
                assert all(r.font.size.pt==13 and r.font.name=='Times New Roman' for p in cell.paragraphs for r in p.runs)
    QA.mkdir(exist_ok=True,parents=True)
    report={'output':str(base.OUTPUT),'native_tables':5,'inline_figures':4,'original_prose_preserved':True,
            'font':'Times New Roman','size_pt':13,'inventory':inventory,
            'evaluation_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest()}
    (QA/'structure-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'native_tables':5,'inline_figures':4,'prose_preserved':True}))


if __name__=='__main__':
    main()

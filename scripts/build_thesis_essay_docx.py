"""Build the prose-only Vietnamese thesis and verify requested Word formatting."""
from pathlib import Path
import json
import re
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/THESIS-DRAFT-CHAPTERS-1-4-VI.md'
OUTPUT = ROOT / 'docs/THESIS-DRAFT-VI.docx'
QA = ROOT / 'reports/thesis-essay-20260910'


def configure_font(style):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(13)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for name in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + name), 'Times New Roman')
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + attr), None)


def main():
    text = SOURCE.read_text(encoding='utf-8')
    assert '\u2014' not in text and ';' not in text
    blocks = [p.strip() for p in text.split('\n\n') if p.strip()]
    doc = Document()
    for borders in list(doc.styles.element.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        configure_font(style)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.widow_control = True
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.space_before = Pt(0)
    normal = doc.styles['Normal'].paragraph_format
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.first_line_indent = Cm(1)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.bold = True
        style.font.italic = False
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(12 if name == 'Heading 1' else 6)
    doc.styles['Title'].paragraph_format.space_before = Pt(0)
    doc.styles['Title'].paragraph_format.space_after = Pt(12)
    expected = []
    for block in blocks:
        if block.startswith('### '):
            content, style = block[4:], 'Heading 2'
        elif block.startswith('## '):
            content, style = block[3:], 'Heading 1'
        elif block.startswith('# '):
            content, style = block[2:], 'Title'
        else:
            content, style = ' '.join(block.splitlines()), 'Normal'
        assert not content.startswith(('- ', '|', '```', '> '))
        p = doc.add_paragraph(content, style)
        expected.append(content)
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), 'vi-VN')
        p.runs[0]._element.get_or_add_rPr().append(lang)
    doc.core_properties.title = expected[0]
    doc.core_properties.subject = 'Khóa luận về hệ thống phân tích ngân hàng có truy nguyên'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    doc.save(OUTPUT)
    check = Document(OUTPUT)
    assert [p.text for p in check.paragraphs] == expected
    assert len(check.tables) == 0 and len(check.inline_shapes) == 0
    assert sum(p.style.name == 'Heading 1' for p in check.paragraphs) == 7
    with ZipFile(OUTPUT) as archive:
        assert not any(re.match(r'word/(header|footer)\d+\.xml', n) for n in archive.namelist())
        assert b'<w:pBdr' not in archive.read('word/styles.xml')
        assert b'<w:pBdr' not in archive.read('word/document.xml')
    for p in check.paragraphs:
        assert p.style.font.name == 'Times New Roman'
        assert p.style.font.size.pt == 13
        assert p.style.paragraph_format.line_spacing == 1.5
        if p.style.name == 'Normal':
            assert p.style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    QA.mkdir(parents=True, exist_ok=True)
    report = {'source': str(SOURCE), 'output': str(OUTPUT), 'paragraphs': len(expected),
              'chapters': 7, 'characters': len(text), 'font': 'Times New Roman',
              'size_pt': 13, 'body_alignment': 'justified', 'line_spacing': 1.5,
              'headers': 0, 'footers': 0, 'tables': 0, 'text_roundtrip': 'passed'}
    (QA / 'structure-check.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))


if __name__ == '__main__':
    main()
